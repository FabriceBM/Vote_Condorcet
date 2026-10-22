"""Concept map of the note: a workbook linking every concept to the bibliography, and every bibliography entry
to the concepts it supports and the key ideas it conveys.

Sources, edited by hand (literature/concept-map/):
    concepts.csv     id, concept, definition and role in the note, note sections
    references.csv   citation key, concept ids (separated by ";"), key ideas conveyed
The bibliographic details come from manuscript/references.bib, and the sections citing each entry are read
from manuscript/sections/*.tex. Each link is entered once, in references.csv; the concept-to-bibliography
sheet is its inversion, so the two directions always agree.

Output: literature/concept-map/concept-map.xlsx, with four sheets:
    About                      purpose, how to read and update, counts and a reciprocity check (formulas)
    Concepts to bibliography   concept | bibliography entries | definition and role | sections | id | count
    Bibliography to concepts   bibliography entry | concepts used | key ideas conveyed | key | cited in | ...
    Matrix                     one row per entry, one column per concept, "x" where they are linked
The counts are COUNTIF formulas on the Matrix sheet; recalculate after writing (see literature/concept-map/README.md).

Run from the repository root:
    python3 tools/concept_map.py            # checks the sources, then writes the workbook
    python3 tools/concept_map.py --check    # checks only
"""
import csv
import os
import re
import sys
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = "literature/concept-map"
BIB = "manuscript/references.bib"
SECTIONS = "manuscript/sections"
OUT = os.path.join(ROOT, "concept-map.xlsx")

FONT = "Arial"
BLUE, LIGHT, GREY = "2A78D6", "EAF2FC", "D9D9D9"
SECTION_LABEL = {"01": "§1 Introduction", "02": "§2 Setting", "03": "§3 Limits", "04": "§4 Veto literature",
                 "05": "§5 Aggregation", "06": "§6 Procedure", "07": "§7 Simulation", "08": "§8 Discussion",
                 "09": "§9 Conclusion", "A": "App. A Proofs", "B": "App. B Algorithms", "C": "App. C Simulation"}

# ---------------------------------------------------------------------------------------------------------
#  Bibliography
# ---------------------------------------------------------------------------------------------------------
ACCENTS = {"'": "\u0301", "`": "\u0300", "^": "\u0302", '"': "\u0308", "~": "\u0303", "c": "\u0327",
           "v": "\u030c"}


def detex(s):
    """LaTeX accents, dashes and braces to plain Unicode text."""
    import unicodedata
    s = re.sub(r"\{\\([cv])\{(\w)\}\}", lambda m: m.group(2) + ACCENTS[m.group(1)], s)
    s = re.sub(r"\\([cv])\{(\w)\}", lambda m: m.group(2) + ACCENTS[m.group(1)], s)
    s = re.sub(r"\{\\(['`^\"~])\{?(\w)\}?\}", lambda m: m.group(2) + ACCENTS[m.group(1)], s)
    s = re.sub(r"\\(['`^\"~])\{?(\w)\}?", lambda m: m.group(2) + ACCENTS[m.group(1)], s)
    s = s.replace("{\\i}", "ı").replace("\\&", "&").replace("\\%", "%").replace("---", "—").replace("--", "–")
    s = s.replace("~", " ")
    s = re.sub(r"\\l?dots(\{\})?\s*", "… ", s)
    s = re.sub(r"\\[a-zA-Z]+\s*", "", s)               # any other command
    s = s.replace("{", "").replace("}", "")
    return unicodedata.normalize("NFC", re.sub(r"\s+", " ", s)).strip()


def read_bib(path):
    """{key: {field: value}} with balanced-brace values; the entry type is in field 'type'."""
    text = open(path, encoding="utf8").read()
    entries = OrderedDict()
    for m in re.finditer(r"@(\w+)\{([^,\s]+),", text):
        kind, key, i = m.group(1).lower(), m.group(2), m.end()
        fields, depth, j = {"type": kind}, 1, i
        while depth and j < len(text):                  # find the end of the entry
            depth += {"{": 1, "}": -1}.get(text[j], 0)
            j += 1
        body = text[i:j - 1]
        for f in re.finditer(r"(\w+)\s*=\s*\{", body):
            k, d, p = f.group(1).lower(), 1, f.end()
            while d:
                d += {"{": 1, "}": -1}.get(body[p], 0)
                p += 1
            fields[k] = body[f.end():p - 1]
        entries[key] = fields
    return entries


def split_authors(authors):
    """[(surname, given names)] from a BibTeX author field; a corporate author (in braces) has no given names."""
    out = []
    for a in re.split(r"\s+and\s+", authors.strip()):
        a = a.strip()
        if a.startswith("{") and a.endswith("}"):       # corporate author, e.g. {Debian Project}
            out.append((detex(a).split(",")[0], ""))
        elif "," in a:
            last, first = a.split(",", 1)
            out.append((detex(last), detex(first)))
        else:
            words = a.split()
            out.append((detex(words[-1]), detex(" ".join(words[:-1]))))
    return out


def surnames(authors):
    return [last for last, _ in split_authors(authors)]


def author_list(authors):
    """'Surname, F. M.' joined by commas, the last with '&'."""
    def initials(first):                               # Jean-François -> J.-F.; Mark Allen -> M. A.
        return " ".join("-".join(p[0] + "." for p in w.split("-") if p) if not w.endswith(".") else w
                        for w in first.split())
    names = [f"{last}, {initials(first)}" if first else last for last, first in split_authors(authors)]
    return names[0] if len(names) == 1 else ", ".join(names[:-1]) + " & " + names[-1]


def short_label(e):
    names = surnames(e.get("author") or e.get("editor", "Anon."))
    who = names[0] if len(names) == 1 else (" & ".join(names) if len(names) == 2 else f"{names[0]} et al.")
    return f"{who} ({detex(e.get('year', 'n.d.'))})"


def full_reference(e):
    authors = author_list(e.get("author") or e.get("editor", ""))
    if not e.get("author") and e.get("editor"):
        authors += " (eds.)"
    parts = [f"{authors} ({detex(e.get('year', 'n.d.'))}). {detex(e.get('title', ''))}."]
    venue = e.get("journal") or e.get("booktitle")
    if venue:
        v = detex(venue)
        if e.get("volume"):
            v += f" {detex(e['volume'])}"
            if e.get("number"):
                v += f"({detex(e['number'])})"
        if e.get("pages"):
            v += f": {detex(e['pages'])}"
        parts.append(v + ".")
    elif e.get("publisher"):
        parts.append(detex(e["publisher"]) + ".")
    elif e.get("howpublished"):
        parts.append(detex(e["howpublished"]) + ".")
    if e.get("note"):
        parts.append(detex(e["note"]) + ".")
    return " ".join(parts)


def link(e):
    if e.get("doi"):
        return "https://doi.org/" + e["doi"].strip()
    if e.get("lccn") and not e.get("url"):
        return f"https://www.loc.gov/item/{e['lccn'].strip()}/"
    return e.get("url", "").strip()


def cited_in():
    """{key: [section labels]} from the \\cite commands of the manuscript sections."""
    out = {}
    for name in sorted(os.listdir(SECTIONS)):
        if not name.endswith(".tex"):
            continue
        label = SECTION_LABEL[name.split("_")[0]]
        text = open(os.path.join(SECTIONS, name), encoding="utf8").read()
        for m in re.finditer(r"\\cite\w*\*?(?:\[[^\]]*\])*\{([^}]+)\}", text):
            for k in m.group(1).split(","):
                out.setdefault(k.strip(), [])
                if label not in out[k.strip()]:
                    out[k.strip()].append(label)
    return out


# ---------------------------------------------------------------------------------------------------------
#  Sources and checks
# ---------------------------------------------------------------------------------------------------------
def load():
    with open(os.path.join(ROOT, "concepts.csv"), encoding="utf8") as fh:
        concepts = OrderedDict((r["id"], r) for r in csv.DictReader(fh))
    with open(os.path.join(ROOT, "references.csv"), encoding="utf8") as fh:
        refs = OrderedDict((r["key"], r) for r in csv.DictReader(fh))
    for r in refs.values():
        r["ids"] = [c.strip() for c in r["concepts"].split(";") if c.strip()]
    return concepts, refs, read_bib(BIB), cited_in()


def check(concepts, refs, bib, cites):
    problems = []
    for k in bib:
        if k not in refs:
            problems.append(f"bibliography entry {k} has no row in references.csv")
    for k, r in refs.items():
        if k not in bib:
            problems.append(f"references.csv row {k} is not in {BIB}")
        if not r["ids"]:
            problems.append(f"{k}: no concept")
        if not r["key_ideas"].strip():
            problems.append(f"{k}: no key ideas")
        for c in r["ids"]:
            if c not in concepts:
                problems.append(f"{k}: unknown concept {c}")
    used = {c for r in refs.values() for c in r["ids"]}
    for c in concepts:
        if c not in used:
            problems.append(f"concept {c} has no bibliography entry")
    for k in cites:
        if k not in bib:
            problems.append(f"the note cites {k}, which is not in {BIB}")
    return problems


# ---------------------------------------------------------------------------------------------------------
#  Workbook
# ---------------------------------------------------------------------------------------------------------
def style_header(ws, row, ncols, height=30):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
    ws.row_dimensions[row].height = height


def body_font(ws, first_row, last_row, ncols):
    thin = Side(style="thin", color=GREY)
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10, bold=cell.font.bold if cell.font else False)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)


def write(concepts, refs, bib, cites):
    order = sorted(refs, key=lambda k: (short_label(bib[k]).lower(), k))        # alphabetical
    cids = list(concepts)
    n_refs, n_con = len(order), len(cids)
    m_first, m_last = 3, 2 + n_refs                                            # Matrix data rows
    m_col0 = 3                                                                 # first concept column
    m_lastcol = get_column_letter(m_col0 + n_con - 1)
    label = {k: f"{short_label(bib[k])}. {detex(bib[k].get('title', ''))}" for k in order}

    wb = Workbook()
    about = wb.active
    about.title = "About"
    ws_c = wb.create_sheet("Concepts to bibliography")
    ws_b = wb.create_sheet("Bibliography to concepts")
    ws_m = wb.create_sheet("Matrix")

    # ---- Matrix -------------------------------------------------------------------------------------
    ws_m["A1"], ws_m["B1"] = "Bibliography entry", "Citation key"
    ws_m["A2"], ws_m["B2"] = "", "Concept id"
    for j, c in enumerate(cids):
        col = m_col0 + j
        ws_m.cell(row=1, column=col, value=concepts[c]["concept"])
        ws_m.cell(row=2, column=col, value=c)
    tot_col = m_col0 + n_con
    ws_m.cell(row=1, column=tot_col, value="Number of concepts")
    for i, k in enumerate(order):
        r = m_first + i
        ws_m.cell(row=r, column=1, value=short_label(bib[k]))
        ws_m.cell(row=r, column=2, value=k)
        for j, c in enumerate(cids):
            if c in refs[k]["ids"]:
                ws_m.cell(row=r, column=m_col0 + j, value="x")
        ws_m.cell(row=r, column=tot_col, value=f'=COUNTIF(C{r}:{m_lastcol}{r},"x")')
    tot_row = m_last + 1
    ws_m.cell(row=tot_row, column=1, value="Number of entries")
    for j in range(n_con):
        L = get_column_letter(m_col0 + j)
        ws_m.cell(row=tot_row, column=m_col0 + j, value=f'=COUNTIF({L}{m_first}:{L}{m_last},"x")')
    ws_m.cell(row=tot_row, column=tot_col, value=f"=SUM({get_column_letter(tot_col)}{m_first}:"
                                                 f"{get_column_letter(tot_col)}{m_last})")
    style_header(ws_m, 1, tot_col, height=190)
    style_header(ws_m, 2, tot_col, height=16)
    for c in range(m_col0, tot_col + 1):
        ws_m.cell(row=1, column=c).alignment = Alignment(text_rotation=90, wrap_text=True, vertical="bottom",
                                                         horizontal="center")
        ws_m.cell(row=2, column=c).alignment = Alignment(horizontal="center")
        ws_m.column_dimensions[get_column_letter(c)].width = 5
    for r in range(m_first, tot_row + 1):
        for c in range(1, tot_col + 1):
            cell = ws_m.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10, bold=(r == tot_row or c == tot_col))
            if c >= m_col0:
                cell.alignment = Alignment(horizontal="center")
                if cell.value == "x":
                    cell.fill = PatternFill("solid", fgColor=LIGHT)
    ws_m.column_dimensions["A"].width = 34
    ws_m.column_dimensions["B"].width = 30
    ws_m.column_dimensions[get_column_letter(tot_col)].width = 9
    ws_m.freeze_panes = "C3"

    # ---- Concepts to bibliography ---------------------------------------------------------------------
    head_c = ["Concept", "Bibliography entries", "Definition and role in the note", "Note sections",
              "Concept id", "Number of entries"]
    for j, h in enumerate(head_c, 1):
        ws_c.cell(row=1, column=j, value=h)
    for i, c in enumerate(cids):
        r = 2 + i
        linked = [k for k in sorted(refs, key=lambda k: (bib[k].get("year", ""), short_label(bib[k])))
                  if c in refs[k]["ids"]]
        ws_c.cell(row=r, column=1, value=concepts[c]["concept"]).font = Font(name=FONT, bold=True)
        ws_c.cell(row=r, column=2, value="\n".join(label[k] for k in linked))
        ws_c.cell(row=r, column=3, value=concepts[c]["definition_and_role"])
        ws_c.cell(row=r, column=4, value=concepts[c]["note_sections"])
        ws_c.cell(row=r, column=5, value=c)
        ws_c.cell(row=r, column=6, value=f'=COUNTIF(INDEX(Matrix!$C${m_first}:${m_lastcol}${m_last},0,'
                                         f'MATCH(E{r},Matrix!$C$2:${m_lastcol}$2,0)),"x")')
    style_header(ws_c, 1, len(head_c))
    body_font(ws_c, 2, 1 + n_con, len(head_c))
    for r in range(2, 2 + n_con):
        ws_c.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)
    for col, w in zip("ABCDEF", (32, 70, 60, 20, 10, 10)):
        ws_c.column_dimensions[col].width = w
    ws_c.freeze_panes = "B2"
    ws_c.auto_filter.ref = f"A1:F{1 + n_con}"

    # ---- Bibliography to concepts ---------------------------------------------------------------------
    head_b = ["Bibliography entry", "Concepts used", "Key ideas conveyed", "Citation key",
              "Cited in (sections of the note)", "Full reference", "DOI or URL", "Still to verify",
              "Number of concepts"]
    for j, h in enumerate(head_b, 1):
        ws_b.cell(row=1, column=j, value=h)
    for i, k in enumerate(order):
        r, e = 2 + i, bib[k]
        ws_b.cell(row=r, column=1, value=label[k])
        ws_b.cell(row=r, column=2, value="\n".join(f"{c} {concepts[c]['concept']}" for c in refs[k]["ids"]))
        ws_b.cell(row=r, column=3, value=refs[k]["key_ideas"])
        ws_b.cell(row=r, column=4, value=k)
        ws_b.cell(row=r, column=5, value=", ".join(cites.get(k, [])) or "not cited yet")
        ws_b.cell(row=r, column=6, value=full_reference(e))
        url = link(e)
        if url:
            cell = ws_b.cell(row=r, column=7, value=url)
            cell.hyperlink = url
        ws_b.cell(row=r, column=8, value=detex(e.get("verify", "")))
        ws_b.cell(row=r, column=9, value=f'=COUNTIF(INDEX(Matrix!$C${m_first}:${m_lastcol}${m_last},'
                                         f'MATCH(D{r},Matrix!$B${m_first}:$B${m_last},0),0),"x")')
    style_header(ws_b, 1, len(head_b))
    body_font(ws_b, 2, 1 + n_refs, len(head_b))
    for r in range(2, 2 + n_refs):
        ws_b.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)
        if ws_b.cell(row=r, column=7).value:
            ws_b.cell(row=r, column=7).font = Font(name=FONT, size=10, color="0563C1", underline="single")
    for col, w in zip("ABCDEFGHI", (40, 40, 60, 24, 22, 60, 30, 30, 10)):
        ws_b.column_dimensions[col].width = w
    ws_b.freeze_panes = "B2"
    ws_b.auto_filter.ref = f"A1:I{1 + n_refs}"

    # ---- About ----------------------------------------------------------------------------------------
    rows = [
        ("Concept map of the note \u201cVoting for a consensus with veto rights\u201d (v0.1)", None),
        ("", None),
        ("What it is", "Every concept of the note linked to the bibliography, and every bibliography entry "
                       "linked to the concepts it supports, with the key ideas it conveys."),
        ("Concepts to bibliography", "One row per concept: the bibliography entries that support it "
                                     "(oldest first), its definition and role in the note, the sections."),
        ("Bibliography to concepts", "One row per entry (alphabetical): the concepts it supports, the key "
                                     "ideas it conveys, where the note cites it, the full reference, the DOI or "
                                     "URL, and any field still to verify."),
        ("Matrix", "One row per entry, one column per concept; x where they are linked. Totals by row and "
                   "column."),
        ("How to update", "Do not edit this workbook: edit literature/concept-map/concepts.csv and "
                          "references.csv (each link is entered once, in references.csv), then run "
                          "python3 tools/concept_map.py from the repository root. Bibliographic details come "
                          "from manuscript/references.bib; the sections citing each entry are read from "
                          "manuscript/sections/."),
        ("", None),
        ("Counts (formulas)", None),
        ("Concepts", f"=COUNTA('Concepts to bibliography'!E2:E{1 + n_con})"),
        ("Bibliography entries", f"=COUNTA('Bibliography to concepts'!D2:D{1 + n_refs})"),
        ("Links, counted by concept", f"=SUM('Concepts to bibliography'!F2:F{1 + n_con})"),
        ("Links, counted by entry", f"=SUM('Bibliography to concepts'!I2:I{1 + n_refs})"),
        ("Links in the matrix", f'=COUNTIF(Matrix!C{m_first}:{m_lastcol}{m_last},"x")'),
        ("Reciprocity check", '=IF(AND(B12=B13,B13=B14),"OK: both directions agree","MISMATCH")'),
    ]
    for i, (a, b) in enumerate(rows, 1):
        about.cell(row=i, column=1, value=a)
        if b is not None:
            about.cell(row=i, column=2, value=b)
    for r in range(1, len(rows) + 1):
        for c in (1, 2):
            cell = about.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10, bold=(c == 1))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    about["A1"].font = Font(name=FONT, size=14, bold=True, color=BLUE)
    about["A9"].font = Font(name=FONT, size=11, bold=True, color=BLUE)
    about.merge_cells("A1:B1")
    about.column_dimensions["A"].width = 28
    about.column_dimensions["B"].width = 100

    for r in range(10, len(rows) + 1):                 # counts: left-aligned next to their labels
        about.cell(row=r, column=2).alignment = Alignment(horizontal="left", vertical="top")
    for ws in wb.worksheets:                           # print: landscape, one page wide
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:2" if ws.title == "Matrix" else "1:1"

    os.makedirs(ROOT, exist_ok=True)
    wb.save(OUT)
    print(f"  {OUT}: {n_con} concepts, {n_refs} entries, "
          f"{sum(len(r['ids']) for r in refs.values())} links")


if __name__ == "__main__":
    concepts, refs, bib, cites = load()
    problems = check(concepts, refs, bib, cites)
    for p in problems:
        print("  PROBLEM:", p)
    if problems:
        sys.exit(1)
    print(f"  sources consistent: {len(concepts)} concepts, {len(refs)} entries, all bibliography keys covered")
    if "--check" not in sys.argv:
        write(concepts, refs, bib, cites)
